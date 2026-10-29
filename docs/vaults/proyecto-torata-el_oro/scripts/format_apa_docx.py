# -*- coding: utf-8 -*-
"""
format_apa_docx.py
==================
Post-procesador OpenXML para asegurar que el documento Word exportado por Quarto
cumpla estrictamente las normas APA 7 / UNL:
  1. Carátula Institucional formal UNL / GAD Torata con logos y metadatos.
  2. Encabezados de página con los tres logos institucionales (UNL, Torata, Economía).
  3. Tipografía Times New Roman estricta en Títulos (Heading 1-4, Title) y Normal.
  4. Páginas horizontales (Landscape) automáticas para tablas anchas (6+ columnas).
  5. Desempaquetado de tablas 1x1 de Quarto/Pandoc:
     - Table Number en negrita (ej. Tabla 1)
     - Table Caption en cursiva debajo del número
     - Bordes APA 7 (solo líneas horizontales, sin líneas verticales)
     - Table Note con 'Nota.' y sangría cero.
"""

import os
import re
import shutil
from copy import deepcopy
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PPR_CHILD_ORDER = [
    'pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widRow',
    'numPr', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
    'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi',
    'adjustRightInd', 'snapToGrid', 'spacing', 'ind', 'contextSpacing',
    'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
    'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange'
]
_PPR_ORDER_MAP = {tag: i for i, tag in enumerate(PPR_CHILD_ORDER)}


def normalize_pPr(pPr):
    children = list(pPr)
    if not children:
        return
    children.sort(key=lambda c: _PPR_ORDER_MAP.get(c.tag.split('}')[-1], 999))
    for c in list(pPr):
        pPr.remove(c)
    for c in children:
        pPr.append(c)


def set_clean_pPr(p_elem, style_id, align=None, keep_next=False):
    for old_pPr in p_elem.findall(qn('w:pPr')):
        p_elem.remove(old_pPr)
    pPr = OxmlElement('w:pPr')
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), style_id)
    pPr.append(pStyle)

    if keep_next:
        pPr.append(OxmlElement('w:keepNext'))

    ind = OxmlElement('w:ind')
    ind.set(qn('w:firstLine'), '0')
    ind.set(qn('w:left'), '0')
    pPr.append(ind)

    if align:
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), align)
        pPr.append(jc)

    normalize_pPr(pPr)
    p_elem.insert(0, pPr)


def set_apa_table_borders(tbl_elem):
    tblPr = tbl_elem.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl_elem.insert(0, tblPr)

    old_ind = tblPr.find(qn('w:tblInd'))
    if old_ind is not None:
        tblPr.remove(old_ind)
    tblInd = OxmlElement('w:tblInd')
    tblInd.set(qn('w:w'), '0')
    tblInd.set(qn('w:type'), 'dxa')
    tblPr.append(tblInd)

    old_w = tblPr.find(qn('w:tblW'))
    if old_w is not None:
        tblPr.remove(old_w)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)

    old_bdr = tblPr.find(qn('w:tblBorders'))
    if old_bdr is not None:
        tblPr.remove(old_bdr)
    tblBorders = OxmlElement('w:tblBorders')
    for b_name, val, sz in [
        ('top', 'single', '8'),
        ('bottom', 'single', '8'),
        ('left', 'none', '0'),
        ('right', 'none', '0'),
        ('insideH', 'none', '0'),
        ('insideV', 'none', '0'),
    ]:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), val)
        b.set(qn('w:sz'), sz)
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        tblBorders.append(b)
    tblPr.append(tblBorders)


def format_inner_table(tbl_elem):
    set_apa_table_borders(tbl_elem)
    rows = tbl_elem.findall(qn('w:tr'))
    for r_idx, tr in enumerate(rows):
        is_header = (r_idx == 0)
        for tc in tr.findall(qn('w:tc')):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = OxmlElement('w:tcPr')
                tc.insert(0, tcPr)
            old_tcBdr = tcPr.find(qn('w:tcBorders'))
            if old_tcBdr is not None:
                tcPr.remove(old_tcBdr)
            if is_header:
                tcBorders = OxmlElement('w:tcBorders')
                top_b = OxmlElement('w:top')
                top_b.set(qn('w:val'), 'single')
                top_b.set(qn('w:sz'), '8')
                top_b.set(qn('w:space'), '0')
                top_b.set(qn('w:color'), '000000')
                bot_b = OxmlElement('w:bottom')
                bot_b.set(qn('w:val'), 'single')
                bot_b.set(qn('w:sz'), '6')
                bot_b.set(qn('w:space'), '0')
                bot_b.set(qn('w:color'), '000000')
                tcBorders.append(top_b)
                tcBorders.append(bot_b)
                tcPr.append(tcBorders)

            for p in tc.findall(qn('w:p')):
                existing_align = 'left'
                for pPr in p.findall(qn('w:pPr')):
                    jc = pPr.find(qn('w:jc'))
                    if jc is not None and jc.get(qn('w:val')):
                        existing_align = jc.get(qn('w:val'))
                style_to_use = 'TableHeader' if is_header else 'TableText'
                set_clean_pPr(p, style_to_use, align=existing_align, keep_next=is_header)
                for r in p.findall(qn('w:r')):
                    rPr = r.find(qn('w:rPr'))
                    if rPr is None:
                        rPr = OxmlElement('w:rPr')
                        r.insert(0, rPr)
                    # Force Times New Roman
                    set_rFonts(rPr, 'Times New Roman')
                    for sz_tag in ['w:sz', 'w:szCs']:
                        old_sz = rPr.find(qn(sz_tag))
                        if old_sz is not None:
                            rPr.remove(old_sz)
                        new_sz = OxmlElement(sz_tag)
                        new_sz.set(qn('w:val'), '20')  # 10 pt
                        rPr.append(new_sz)


def set_rFonts(rPr, font_name='Times New Roman'):
    old_rf = rPr.find(qn('w:rFonts'))
    if old_rf is not None:
        rPr.remove(old_rf)
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), font_name)
    rf.set(qn('w:hAnsi'), font_name)
    rf.set(qn('w:cs'), font_name)
    rf.set(qn('w:eastAsia'), font_name)
    rPr.insert(0, rf)


def make_text_para(text, style_id, bold=False, italic=False, keep_next=True):
    p = OxmlElement('w:p')
    set_clean_pPr(p, style_id, align='left', keep_next=keep_next)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    set_rFonts(rPr, 'Times New Roman')
    if bold:
        rPr.append(OxmlElement('w:b'))
        rPr.append(OxmlElement('w:bCs'))
    if italic:
        rPr.append(OxmlElement('w:i'))
        rPr.append(OxmlElement('w:iCs'))
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '20')
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), '20')
    rPr.append(sz)
    rPr.append(szCs)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    r.append(t)
    p.append(r)
    return p


def make_sectPr(landscape=False):
    sectPr = OxmlElement('w:sectPr')
    pgSz = OxmlElement('w:pgSz')
    if landscape:
        pgSz.set(qn('w:w'), '15840')  # 11 in
        pgSz.set(qn('w:h'), '12240')  # 8.5 in
        pgSz.set(qn('w:orient'), 'landscape')
    else:
        pgSz.set(qn('w:w'), '12240')  # 8.5 in
        pgSz.set(qn('w:h'), '15840')  # 11 in
        pgSz.set(qn('w:orient'), 'portrait')
    sectPr.append(pgSz)

    pgMar = OxmlElement('w:pgMar')
    pgMar.set(qn('w:top'), '1440')     # 1 in = 25.4 mm
    pgMar.set(qn('w:bottom'), '1440')
    pgMar.set(qn('w:left'), '1440')
    pgMar.set(qn('w:right'), '1440')
    sectPr.append(pgMar)
    return sectPr


def apply_header_logos(doc, assets_dir: Path):
    logo_unl = assets_dir / 'logo_unl.png'
    logo_torata = assets_dir / 'logo_torata_actualizado.png'
    logo_eco = assets_dir / 'logo_economia.jpg'

    # Section 0 is the master section for headers
    sec0 = doc.sections[0]
    sec0.different_first_page_header_footer = True

    # Ensure first page header is empty (cover page)
    first_h = sec0.first_page_header
    for p in list(first_h.paragraphs):
        p.text = ''
    for t in list(first_h.tables):
        first_h._element.remove(t._element)

    # Master header for normal pages: remove old tables
    header = sec0.header
    for p in list(header.paragraphs):
        p.text = ''
    for t in list(header.tables):
        header._element.remove(t._element)

    tbl = header.add_table(rows=1, cols=3, width=Inches(6.5))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._element.xpath('w:tblPr')[0]
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + b)
        el.set(qn('w:val'), 'none')
        tblBorders.append(el)
    tblPr.append(tblBorders)

    cell_l, cell_c, cell_r = tbl.rows[0].cells
    cell_l.width = Inches(2.0)
    cell_c.width = Inches(2.5)
    cell_r.width = Inches(2.0)

    # UNL logo left
    if logo_unl.exists():
        pl = cell_l.paragraphs[0]
        pl.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pl.add_run().add_picture(str(logo_unl), height=Inches(0.60))

    # Torata logo center
    if logo_torata.exists():
        pc = cell_c.paragraphs[0]
        pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pc.add_run().add_picture(str(logo_torata), height=Inches(0.60))

    # Economia logo right
    if logo_eco.exists():
        pr = cell_r.paragraphs[0]
        pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pr.add_run().add_picture(str(logo_eco), height=Inches(0.60))

    # For subsequent sections, keep them linked to previous so they inherit the exact header
    for sec in doc.sections[1:]:
        sec.different_first_page_header_footer = False
        sec.header.is_linked_to_previous = True


def insert_cover_page(doc, assets_dir: Path):
    """
    Inserta la carátula institucional UNL / GAD Torata al inicio del documento.
    """
    body = doc._element.body
    logo_unl = assets_dir / 'logo_unl.png'
    logo_torata = assets_dir / 'logo_torata_actualizado.png'
    logo_eco = assets_dir / 'logo_economia.jpg'

    # Remove default Pandoc Title and Author paragraphs from document start
    for p in list(body.findall(qn('w:p'))):
        pPr = p.find(qn('w:pPr'))
        if pPr is not None:
            pStyle = pPr.find(qn('w:pStyle'))
            if pStyle is not None and pStyle.get(qn('w:val')) in ['Title', 'Author', 'Date']:
                body.remove(p)
                continue
        # Stop once we reach TOC or Heading 1
        break

    # Build cover elements container
    cover_elements = []

    def p_cover(text="", font_size=12, bold=False, italic=False, align="center", space_before=0, space_after=6):
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), align)
        pPr.append(jc)
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:before'), str(int(space_before * 20)))
        sp.set(qn('w:after'), str(int(space_after * 20)))
        sp.set(qn('w:line'), '276')  # 1.15 line spacing
        sp.set(qn('w:lineRule'), 'auto')
        pPr.append(sp)
        normalize_pPr(pPr)
        p.append(pPr)

        if text:
            r = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')
            set_rFonts(rPr, 'Times New Roman')
            if bold:
                rPr.append(OxmlElement('w:b'))
                rPr.append(OxmlElement('w:bCs'))
            if italic:
                rPr.append(OxmlElement('w:i'))
                rPr.append(OxmlElement('w:iCs'))
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), str(int(font_size * 2)))
            szCs = OxmlElement('w:szCs')
            szCs.set(qn('w:val'), str(int(font_size * 2)))
            rPr.append(sz)
            rPr.append(szCs)
            r.append(rPr)
            t = OxmlElement('w:t')
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            t.text = text
            r.append(t)
            p.append(r)
        return p

    # 1. Logos table for Cover Page
    dummy_doc = docx.Document()
    tbl = dummy_doc.add_table(rows=1, cols=3)
    tblPr = tbl._element.xpath('w:tblPr')[0]
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + b)
        el.set(qn('w:val'), 'none')
        tblBorders.append(el)
    tblPr.append(tblBorders)

    c_l, c_c, c_r = tbl.rows[0].cells
    c_l.width = Inches(2.0)
    c_c.width = Inches(2.5)
    c_r.width = Inches(2.0)

    if logo_unl.exists():
        c_l.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        c_l.paragraphs[0].add_run().add_picture(str(logo_unl), height=Inches(0.95))
    if logo_torata.exists():
        c_c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_c.paragraphs[0].add_run().add_picture(str(logo_torata), height=Inches(0.95))
    if logo_eco.exists():
        c_r.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        c_r.paragraphs[0].add_run().add_picture(str(logo_eco), height=Inches(0.95))

    cover_elements.append(tbl._element)
    cover_elements.append(p_cover("", space_after=12))

    # 2. Text elements of Cover
    cover_elements.append(p_cover("UNIVERSIDAD NACIONAL DE LOJA", font_size=14, bold=True, space_after=4))
    cover_elements.append(p_cover("FACULTAD JURÍDICA, SOCIAL Y ADMINISTRATIVA", font_size=12, bold=True, space_after=4))
    cover_elements.append(p_cover("CARRERA DE ECONOMÍA", font_size=12, bold=True, space_after=18))

    cover_elements.append(p_cover("PROYECTO DE FACTIBILIDAD Y PLAN DE INVERSIÓN PÚBLICA:", font_size=11, bold=True, italic=True, space_after=6))
    cover_elements.append(p_cover("CONECTIVIDAD VIAL Y CONSTRUCCIÓN DEL PUENTE VEHICULAR SOBRE EL RÍO SANTA ROSA EN EL SECTOR EL GUAYABO - LA CHILCA, PARROQUIA RURAL TORATA",
                                  font_size=13, bold=True, space_after=20))

    cover_elements.append(p_cover("PROYECTO DE INVERSIÓN PÚBLICA - SÉPTIMO CICLO", font_size=11, bold=True, space_after=18))

    cover_elements.append(p_cover("AUTORES:", font_size=11, bold=True, space_after=4))
    for autor in [
        "David Andrés Pérez Castillo",
        "Luis Gustavo Hurtado Cuenca",
        "Erick Fabricio Condoy Seraquive",
        "Lizeth Camila Correa Sánchez"
    ]:
        cover_elements.append(p_cover(autor, font_size=11, bold=False, space_after=2))

    cover_elements.append(p_cover("", space_after=12))
    cover_elements.append(p_cover("DOCENTE:", font_size=11, bold=True, space_after=2))
    cover_elements.append(p_cover("Econ. Jose Job Chamba Tandazo", font_size=11, space_after=18))

    cover_elements.append(p_cover("LOJA - ECUADOR", font_size=11, bold=True, space_after=2))
    cover_elements.append(p_cover("Octubre de 2026", font_size=11, space_after=12))

    # Page break after cover
    p_break = p_cover()
    r_br = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r_br.append(br)
    p_break.append(r_br)
    cover_elements.append(p_break)

    # Insert cover elements at the very top of body
    for elem in reversed(cover_elements):
        body.insert(0, elem)


def enforce_times_new_roman_headings(doc):
    """
    Asegura que todos los párrafos de títulos y texto tengan Times New Roman explícito.
    """
    body = doc._element.body
    for p in body.findall(qn('w:p')):
        pPr = p.find(qn('w:pPr'))
        is_heading = False
        if pPr is not None:
            pStyle = pPr.find(qn('w:pStyle'))
            if pStyle is not None:
                sval = pStyle.get(qn('w:val'), '')
                if 'Heading' in sval or sval in ['Title', 'Subtitle', 'TOCHeading']:
                    is_heading = True

        for r in p.findall(qn('w:r')):
            rPr = r.find(qn('w:rPr'))
            if rPr is None:
                rPr = OxmlElement('w:rPr')
                r.insert(0, rPr)
            set_rFonts(rPr, 'Times New Roman')
            if is_heading:
                # Ensure APA bold/format
                if rPr.find(qn('w:b')) is None:
                    rPr.append(OxmlElement('w:b'))
                    rPr.append(OxmlElement('w:bCs'))


def convert_wide_tables_to_landscape(doc):
    """
    Convierte dinámicamente las secciones de tablas anchas (6 o más columnas) a formato horizontal.
    """
    body = doc._element.body
    for tbl in list(body.findall(qn('w:tbl'))):
        rows = tbl.findall(qn('w:tr'))
        if not rows:
            continue
        first_row_cells = rows[0].findall(qn('w:tc'))
        cols_count = len(first_row_cells)

        # Check if table is wide or landscape
        if cols_count >= 6:
            # Find the caption paragraph(s) before the table
            p_first = tbl
            p_scan = tbl.getprevious()
            while p_scan is not None and p_scan.tag.endswith('p'):
                txt = ''.join(x.text or '' for x in p_scan.findall('.//' + qn('w:t'))).strip()
                if txt.startswith('Tabla') or txt.startswith('Figura'):
                    p_first = p_scan
                elif txt == '' and p_first != tbl:
                    pass
                else:
                    break
                p_scan = p_scan.getprevious()

            # Find the note paragraph after the table
            p_last = tbl
            p_scan_after = tbl.getnext()
            while p_scan_after is not None and p_scan_after.tag.endswith('p'):
                txt = ''.join(x.text or '' for x in p_scan_after.findall('.//' + qn('w:t'))).strip()
                if txt.startswith('Nota.') or txt == '':
                    p_last = p_scan_after
                    if txt.startswith('Nota.'):
                        break
                else:
                    break
                p_scan_after = p_scan_after.getnext()

            # Check if previous element already has a landscape break
            prev_elem = p_first.getprevious()
            has_break_before = False
            if prev_elem is not None and prev_elem.tag.endswith('p'):
                pPr = prev_elem.find(qn('w:pPr'))
                if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
                    has_break_before = True

            if not has_break_before:
                p_break_before = OxmlElement('w:p')
                pPr_b = OxmlElement('w:pPr')
                pPr_b.append(make_sectPr(landscape=False))
                p_break_before.append(pPr_b)
                p_first.addprevious(p_break_before)

            # Insert landscape section break after table note
            p_break_after = OxmlElement('w:p')
            pPr_a = OxmlElement('w:pPr')
            pPr_a.append(make_sectPr(landscape=True))
            p_break_after.append(pPr_a)
            p_last.addnext(p_break_after)


def process_docx(docx_path: Path):
    doc = docx.Document(str(docx_path))
    body = doc._element.body
    vault_dir = docx_path.parent
    assets_dir = vault_dir / 'asset'
    if not assets_dir.exists():
        assets_dir = vault_dir.parent.parent / 'asset'

    # 1. Unwrap Quarto 1x1 wrapper tables (both Table wrappers and Figure wrappers)
    for tbl in list(body.findall(qn('w:tbl'))):
        rows = tbl.findall(qn('w:tr'))
        if len(rows) == 1:
            cells = rows[0].findall(qn('w:tc'))
            if len(cells) == 1:
                tc = cells[0]
                paras = tc.findall(qn('w:p'))
                inner_tbls = tc.findall(qn('w:tbl'))
                if paras:
                    first_p_text = ''.join(t.text or '' for t in paras[0].findall('.//' + qn('w:t'))).strip()
                    m = re.match(r'^((?:Tabla|Figura)\s+\d+)\.\s*(.+)$', first_p_text, flags=re.DOTALL)
                    if m:
                        label_num = m.group(1).replace('\xa0', ' ')
                        title_str = m.group(2).strip()
                        is_table = label_num.startswith('Tabla')

                        p_num = make_text_para(label_num, 'TableNumber', bold=True, italic=False, keep_next=True)
                        p_cap = make_text_para(title_str, 'TableCaption' if is_table else 'ImageCaption',
                                               bold=False, italic=True, keep_next=True)

                        tbl.addprevious(p_num)
                        tbl.addprevious(p_cap)

                        if inner_tbls:
                            inner_tbl = deepcopy(inner_tbls[0])
                            format_inner_table(inner_tbl)
                            tbl.addprevious(inner_tbl)
                        else:
                            for fig_p in paras[1:]:
                                fig_p_copy = deepcopy(fig_p)
                                set_clean_pPr(fig_p_copy, 'CaptionedFigure', align='center', keep_next=False)
                                tbl.addprevious(fig_p_copy)

                        body.remove(tbl)

    # 2. Standalone captions and Table Note formatting
    for p in list(body.findall(qn('w:p'))):
        txt = ''.join(t.text or '' for t in p.findall('.//' + qn('w:t'))).strip()
        m_fig = re.match(r'^((?:Tabla|Figura)\s+\d+)\.\s*(.+)$', txt, flags=re.DOTALL)
        if m_fig:
            label_num = m_fig.group(1).replace('\xa0', ' ')
            title_str = m_fig.group(2).strip()
            p_num = make_text_para(label_num, 'TableNumber', bold=True, italic=False, keep_next=True)
            p_cap = make_text_para(title_str, 'ImageCaption', bold=False, italic=True, keep_next=True)
            p.addprevious(p_num)
            p.addprevious(p_cap)
            body.remove(p)
            continue

        if txt.startswith('Nota.'):
            set_clean_pPr(p, 'TableNote', align='both', keep_next=False)
            for r in p.findall(qn('w:r')):
                rPr = r.find(qn('w:rPr'))
                if rPr is None:
                    rPr = OxmlElement('w:rPr')
                    r.insert(0, rPr)
                set_rFonts(rPr, 'Times New Roman')
                for sz_tag in ['w:sz', 'w:szCs']:
                    old_sz = rPr.find(qn(sz_tag))
                    if old_sz is not None:
                        rPr.remove(old_sz)
                    new_sz = OxmlElement(sz_tag)
                    new_sz.set(qn('w:val'), '20')
                    rPr.append(new_sz)

    # 3. Horizontal orientation for wide tables
    convert_wide_tables_to_landscape(doc)

    # 4. Enforce Times New Roman on all headings
    enforce_times_new_roman_headings(doc)

    # 5. Insert Cover Page
    insert_cover_page(doc, assets_dir)

    # 6. Apply 3 Logos to Header
    apply_header_logos(doc, assets_dir)

    # Normalize all pPr
    for pPr in body.iter(qn('w:pPr')):
        normalize_pPr(pPr)

    doc.save(str(docx_path))
    print(f'Formatted APA 7 styles, cover, landscape and headers in: {docx_path}')


if __name__ == '__main__':
    site_docx = Path(r'd:\07_Ciclo\septimo_ciclo\public_investment_projects\_site\docs\vaults\proyecto-torata-el_oro\index.docx')
    vault_docx = Path(r'd:\07_Ciclo\septimo_ciclo\public_investment_projects\docs\vaults\proyecto-torata-el_oro\index.docx')
    if site_docx.exists():
        process_docx(site_docx)
        shutil.copy2(site_docx, vault_docx)
        print(f'Synced to: {vault_docx}')
    elif vault_docx.exists():
        process_docx(vault_docx)
