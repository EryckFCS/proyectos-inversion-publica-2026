import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Márgenes estándar
for s in doc.sections:
    s.top_margin = Inches(1.0)
    s.bottom_margin = Inches(1.0)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(1.0)

# Encabezado institucional
h1 = doc.add_heading('UNIVERSIDAD NACIONAL DE LOJA', 1)
h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
h2 = doc.add_heading('FACULTAD JURÍDICA, SOCIAL Y ADMINISTRATIVA | CARRERA DE ECONOMÍA', 2)
h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub = doc.add_paragraph('Proyectos de Inversión Pública - Séptimo Ciclo | Docente: Econ. José Job Chamba Tandazo, Mg. Sc.')
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

h3 = doc.add_heading('GUÍA DE ENTREVISTA A INFORMANTES CLAVE Y AUTORIDADES LOCALES\nDIAGNÓSTICO TERRITORIAL Y LÍNEA BASE PARA PROYECTO DE INVERSIÓN PÚBLICA:\n"CONSTRUCCIÓN / REPOSICIÓN DE PUENTE EN LA PARROQUIA TORATA, EL ORO"', 3)
h3.alignment = WD_ALIGN_PARAGRAPH.CENTER

p_prot = doc.add_paragraph(
    'Buenos días/tardes. Somos estudiantes e investigadores de la Carrera de Economía de la Universidad Nacional de Loja. '
    'Estamos realizando el levantamiento de información de campo y diagnóstico integral del territorio para sustentar '
    'la viabilidad técnica, social, económica y ambiental del proyecto de inversión pública enfocado en la '
    'Construcción y Conectividad mediante Infraestructura de Puentes en la Parroquia Rural Torata, Cantón Santa Rosa, '
    'Provincia de El Oro. Agradecemos su colaboración respondiendo a la presente entrevista estructurada.'
)

# Módulo 0: Ficha Técnica de la Entrevista
doc.add_heading('MÓDULO 0: DATOS DE CONTROL Y LOCALIZACIÓN DE LA ENTREVISTA', 2)
t0 = doc.add_table(rows=4, cols=3)
t0.style = 'Table Grid'
t0.alignment = WD_TABLE_ALIGNMENT.CENTER

t0.cell(0, 0).text = 'Provincia / Cantón / Parroquia:'
t0.cell(0, 1).text = 'El Oro / Santa Rosa / Torata'
t0.cell(0, 2).text = 'Fecha: ____ / ____ / 2026'

t0.cell(1, 0).text = 'Lugar / Sector:'
t0.cell(1, 1).text = 'Sede GAD Parroquial Rural de Torata / Sector del Puente'
t0.cell(1, 2).text = 'Hora inicio: ____:____'

t0.cell(2, 0).text = 'Entrevistado(a) / Cargo:'
t0.cell(2, 1).text = 'Autoridad Local / Representante Comunitario / Líder Parroquial'
t0.cell(2, 2).text = 'Hora fin: ____:____'

t0.cell(3, 0).text = 'Equipo Evaluador UNL:'
t0.cell(3, 1).text = 'Estudiantes - Carrera de Economía'
t0.cell(3, 2).text = 'Coordenadas GPS (WGS84):\nLat: __________ Long: __________'

# MÓDULO I: CARACTERIZACIÓN GENERAL Y PLANIFICACIÓN
doc.add_heading('MÓDULO I: CARACTERIZACIÓN TERRITORIAL Y MARCO DE PLANIFICACIÓN', 2)

doc.add_paragraph('1.1 Marco Político-Administrativo y Planificación Institucional:')
doc.add_paragraph(
    '1.1.1 Situación del Plan de Desarrollo y Ordenamiento Territorial (PDOT) de Torata:\n'
    '[ ] Vigente y aprobado (2023-2027)      [ ] En proceso de actualización      [ ] Desactualizado\n'
    'Observaciones / Articulación con GAD Municipal de Santa Rosa y Prefectura de El Oro:\n'
    '_____________________________________________________________________________________________________'
)

doc.add_paragraph(
    '1.1.2 Comunidades, recintos y sectores prioritarios atendidos por el GAD:\n'
    '¿Cuáles son las comunidades o barrios rurales con mayor necesidad de integración vial e infraestructura física?\n'
    '_____________________________________________________________________________________________________'
)

doc.add_paragraph('1.2 Características Biogeográficas y Entorno Hídrico:')
t_bio = doc.add_table(rows=5, cols=3)
t_bio.style = 'Table Grid'
t_bio.alignment = WD_TABLE_ALIGNMENT.CENTER
t_bio.cell(0, 0).text = 'Variable Territorial'
t_bio.cell(0, 1).text = 'Detalle en la Parroquia Torata'
t_bio.cell(0, 2).text = 'Registro / Estimación de Campo'
t_bio.cell(1, 0).text = 'Ríos y Afluentes Principales'
t_bio.cell(1, 1).text = 'Río Torata y microcuencas de drenaje'
t_bio.cell(1, 2).text = 'Cuenca hidrográfica'
t_bio.cell(2, 0).text = 'Régimen de Lluvias e Inviernos'
t_bio.cell(2, 1).text = 'Comportamiento de crecidas en temporada invernal (Ene - May)'
t_bio.cell(2, 2).text = 'Nivel de riesgo por crecidas'
t_bio.cell(3, 0).text = 'Relieve y Topografía'
t_bio.cell(3, 1).text = 'Transición estribación andina - zona baja / piedemonte'
t_bio.cell(3, 2).text = 'Inclinación / estabilidad taludes'
t_bio.cell(4, 0).text = 'Clima y Temperatura Media'
t_bio.cell(4, 1).text = 'Subtropical / Cálido húmedo (22 °C - 27 °C)'
t_bio.cell(4, 2).text = '°C / Humedad'

# MÓDULO II: DIAGNÓSTICO ESPECÍFICO DEL PUENTE Y CONECTIVIDAD VIAL
doc.add_heading('MÓDULO II: DIAGNÓSTICO ESPECÍFICO DEL PUENTE Y SITUACIÓN VIAL ACTUAL', 2)

doc.add_paragraph(
    '2.1 Identificación y Estado Actual del Paso / Puente Objeto del Proyecto:\n'
    '2.1.1 ¿Cuál es la situación actual del cruce o infraestructura de puente en el sector prioritario de Torata?\n'
    '[ ] Inexistente (cruce por badén de tierra / vado natural en el cauce del río)\n'
    '[ ] Estructura colapsada o destruida por crecidas invernales anteriores\n'
    '[ ] Puente provisional / artesanal de madera o caña con alto riesgo de caída\n'
    '[ ] Puente antiguo metálico / de hormigón en estado crítico (deterioro estructural, socavación)\n'
    'Detalle técnico o antecedentes del cruce:\n'
    '_____________________________________________________________________________________________________'
)

doc.add_paragraph(
    '2.1.2 ¿Qué comunidades, recintos, sectores o fincas quedan incomunicadas cuando el río crece?\n'
    'Recintos afectados: _________________________________________________________________________________\n'
    'Población directamente afectada estimada: ________________ personas / ________________ familias.'
)

doc.add_paragraph(
    '2.1.3 ¿Cuántos días o semanas al año se interrumpe totalmente el paso de vehículos y peatones debido al río?\n'
    '[ ] Menos de 15 días acumulados al año\n'
    '[ ] Entre 15 y 45 días al año\n'
    '[ ] Más de 45 días al año (frecuente incomunicación durante la estación lluviosa)\n'
    '[ ] Paso permanentemente riesgoso y precario todo el año'
)

doc.add_paragraph(
    '2.1.4 Tipología y características del transporte que requiere cruzar el puente:\n'
    '[ ] Peatonal y acémilas (caballos/mulas)\n'
    '[ ] Motocicletas y vehículos livianos (camionetas familiares)\n'
    '[ ] Camiones de carga mediana (transporte de cosechas agrícolas y víveres)\n'
    '[ ] Transporte pesado / camiones de ganado o material de construcción\n'
    '[ ] Buses de transporte escolar o público interparroquial'
)

# MÓDULO III: IMPACTO SOCIOECONÓMICO Y PRODUCTIVO
doc.add_heading('MÓDULO III: COMPONENTE PRODUCTIVO, SOCIAL Y PÉRDIDAS ECONÓMICAS', 2)

doc.add_paragraph(
    '3.1 Cadenas Productivas Afectadas por la Falta del Puente:\n'
    '¿Cuáles son los principales productos de la zona cuya comercialización se ve perjudicada por la falta de un puente seguro?\n'
    '[ ] Cacao nacional y fino de aroma\n'
    '[ ] Plátano / Guineo / Frutales\n'
    '[ ] Maíz y granos básicos\n'
    '[ ] Ganadería bovina (leche y carne) y derivados\n'
    '[ ] Caña de azúcar / Panela\n'
    '[ ] Cítricos (naranja, mandarina, limón)\n'
    '[ ] Otros: _________________________________________________________________________________________'
)

doc.add_paragraph(
    '3.2 Pérdidas Económicas y Costos de Transporte Adicionales:\n'
    '3.2.1 ¿Qué ocurre con la producción agropecuaria cuando el río no da paso?\n'
    '[ ] Se pierde o pudre la cosecha en las fincas por no poder sacarla a tiempo\n'
    '[ ] Se vende a intermediarios a precios castigados o de remate\n'
    '[ ] Se debe pagar trasbordo manual o acémilas a costos excesivos (sobrecosto de flete)\n'
    '3.2.2 Estimación de sobrecosto por flete o trasbordo para los productores:\n'
    '¿Cuánto incrementa el costo de transporte por bulto/gaveta/quintal por no tener puente? $ ______________'
)

doc.add_paragraph(
    '3.3 Acceso a Servicios Esenciales de Salud y Educación:\n'
    '3.3.1 ¿Cómo afecta la falta del puente al acceso escolar de niños y jóvenes de los recintos hacia Torata o Santa Rosa?\n'
    '[ ] Los niños pierden clases repetidamente durante el invierno por peligro de cruzar\n'
    '[ ] Cruce a pie arriesgando la vida de estudiantes y padres de familia\n'
    '[ ] No existe afectación relevante en educación\n'
    '3.3.2 En caso de emergencias médicas graves o traslados al Centro de Salud o al Hospital de Santa Rosa:\n'
    '[ ] Imposibilidad total de ingreso de ambulancias o vehículos de emergencia\n'
    '[ ] Necesidad de trasladar pacientes en hamacas / a pie a través del río crecido\n'
    '[ ] Tiempo de retraso crítico en la atención médica: ________________ horas de demora adicional.'
)

# MÓDULO IV: ANÁLISIS DE RIESGOS, AMBIENTE Y GESTIÓN INSTITUCIONAL
doc.add_heading('MÓDULO IV: RIESGOS NATURALES, SEGURIDAD Y GESTIÓN PÚBLICA', 2)

doc.add_paragraph(
    '4.1 Seguridad Ciudadana y Antecedentes de Accidentes:\n'
    '¿Se han registrado siniestros, pérdidas materiales (vehículos arrastrados) o accidentes personales en este paso?\n'
    '[ ] Sí       [ ] No\n'
    'Describa antecedentes o eventos críticos registrados:\n'
    '_____________________________________________________________________________________________________\n'
    '_____________________________________________________________________________________________________'
)

doc.add_paragraph(
    '4.2 Competencias Institucionales y Gestiones Previas:\n'
    '4.2.1 ¿Qué gestiones se han realizado ante los diferentes niveles de gobierno para la construcción de este puente?\n'
    '- Prefectura de El Oro (Competencia de vialidad rural y puentes): [ ] Solicitud ingresada  [ ] Proyecto/Estudios  [ ] Sin respuesta\n'
    '- GAD Municipal de Santa Rosa:                                    [ ] Convenio concurrente [ ] En trámite         [ ] Sin intervención\n'
    '- GAD Parroquial Rural de Torata:                                 [ ] Priorizado en POA    [ ] Gestión directa    [ ] Presupuesto insuficiente\n'
    'Detalle de acuerdos o estado de estudios técnicos previos (topografía, mecánica de suelos, hidrología):\n'
    '_____________________________________________________________________________________________________'
)

doc.add_paragraph(
    '4.3 Priorización y Criterio Técnico del Proyecto Formulador UNL:\n'
    '4.3.1 ¿En qué escala prioriza la comunidad y el GAD Parroquial la construcción definitiva del puente?\n'
    '[ ] Prioridad Máxima (Urgencia de vida e integración productiva)\n'
    '[ ] Prioridad Alta\n'
    '[ ] Prioridad Media\n'
    '4.3.2 Criterio sobre el apoyo técnico y formulación académica de la Universidad Nacional de Loja:\n'
    '_____________________________________________________________________________________________________\n'
    '_____________________________________________________________________________________________________'
)

doc.add_paragraph(
    'Firmas de Responsabilidad:\n\n\n'
    '_____________________________________                  _____________________________________\n'
    '       Entrevistado / Autoridad                                Investigador UNL Evaluador\n'
    '       GAD Parroquial Torata                                   Carrera de Economía'
)

out_doc = r'd:\07_Ciclo\septimo_ciclo\public_investment_projects\docs\vaults\proyecto-saracay\levantamiento_de_informacion\Guia_Entrevista_Diagnostico_Torata_Puente_ElOro.docx'
doc.save(out_doc)
print('Documento Word de Entrevista Torata generado exitosamente en:', out_doc)
