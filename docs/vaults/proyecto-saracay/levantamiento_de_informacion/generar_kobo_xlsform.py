import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Setup sheets
ws_survey = wb.active
ws_survey.title = "survey"
ws_choices = wb.create_sheet(title="choices")
ws_settings = wb.create_sheet(title="settings")

# Header styling
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")

# 1. SETTINGS SHEET
ws_settings.append(["form_title", "form_id", "version", "default_language"])
ws_settings.append([
    "Diagnóstico Territorial Saracay 2026 (Piñas - El Oro) - UNL",
    "unl_diagnostico_saracay_eloro_2026",
    "2026101501",
    "Spanish (es)"
])

# 2. CHOICES SHEET
choices_headers = ["list_name", "name", "label"]
ws_choices.append(choices_headers)

choices_data = [
    # si_no
    ("si_no", "si", "Sí"),
    ("si_no", "no", "No"),
    
    # si_no_ns
    ("si_no_ns", "si", "Sí"),
    ("si_no_ns", "no", "No"),
    ("si_no_ns", "ns_nr", "No sabe / No responde"),
    
    # cargo_rol
    ("cargo_rol", "presidente_gad", "Presidente del GAD Parroquial (Sr. Jorge Homero Ochoa Sánchez)"),
    ("cargo_rol", "vocal_gad", "Vocal del GAD Parroquial"),
    ("cargo_rol", "lider_comunitario", "Líder / Representante Comunitario / Recinto"),
    ("cargo_rol", "personal_salud", "Personal de Salud (Centro de Salud Saracay)"),
    ("cargo_rol", "directivo_educativo", "Directivo / Docente Unidad Educativa Saracay"),
    ("cargo_rol", "productor_agropecuario", "Productor agropecuario (Caña, Maíz, Ganadería, Cacao)"),
    ("cargo_rol", "comerciante_servicios", "Comerciante / Gastronómico / Transportista en La Y"),
    ("cargo_rol", "morador", "Habitante / Morador de la parroquia"),
    ("cargo_rol", "otro", "Otro (especificar)"),
    
    # rango_edad
    ("rango_edad", "18_29", "18 a 29 años"),
    ("rango_edad", "30_45", "30 a 45 años"),
    ("rango_edad", "46_60", "46 a 60 años"),
    ("rango_edad", "mas_60", "Más de 60 años"),
    
    # sitio_sector (Ciudadelas y recintos de Saracay, Piñas, El Oro)
    ("sitio_sector", "cabecera_central", "Cabecera Parroquial - Ciudadela Central"),
    ("sitio_sector", "sector_la_y", "Sector La Y de Saracay (Nodo vial E585 / E25)"),
    ("sitio_sector", "ciud_el_bosque", "Ciudadela El Bosque"),
    ("sitio_sector", "ciud_el_tamarindo", "Ciudadela El Tamarindo"),
    ("sitio_sector", "ciud_la_inmaculada", "Ciudadela La Inmaculada"),
    ("sitio_sector", "ciud_las_brisas", "Ciudadela Las Brisas"),
    ("sitio_sector", "recinto_platanillos", "Recinto Platanillos"),
    ("sitio_sector", "recinto_panupali", "Recinto Panupali"),
    ("sitio_sector", "recinto_naranjos", "Recinto Naranjos"),
    ("sitio_sector", "recinto_san_jose", "Recinto San José"),
    ("sitio_sector", "recinto_fatima", "Recinto Fátima"),
    ("sitio_sector", "recinto_damas", "Recinto Damas"),
    ("sitio_sector", "otro_recinto", "Otro recinto o comunidad"),
    
    # calidad_agua
    ("calidad_agua", "buena_potable", "Buena (tratada / potable apta para consumo)"),
    ("calidad_agua", "regular_entubada", "Regular (agua entubada sin cloración continua)"),
    ("calidad_agua", "pozo_vertiente", "Pozo somero / vertiente o quebrada particular"),
    ("calidad_agua", "tanquero_compra", "Abastecimiento por tanquero o compra en bidón"),
    ("calidad_agua", "no_dispone", "Sin acceso a red formal"),
    
    # continuidad_agua
    ("continuidad_agua", "continua_24h", "Continua las 24 horas del día"),
    ("continuidad_agua", "parcial_12_18h", "Parcial (12 a 18 horas diarias)"),
    ("continuidad_agua", "intermitente_por_turnos", "Intermitente (pocas horas al día o por turnos)"),
    ("continuidad_agua", "critica_cortes", "Crítica (cortes prolongados o días alternos)"),
    
    # saneamiento_vivienda
    ("saneamiento_vivienda", "red_alcantarillado", "Red pública de alcantarillado sanitario"),
    ("saneamiento_vivienda", "pozo_septico", "Pozo séptico individual adecuado"),
    ("saneamiento_vivienda", "pozo_ciego_letrina", "Pozo ciego / letrina"),
    ("saneamiento_vivienda", "descarga_canal_rio", "Descarga directa al río Saracay o quebradas"),
    
    # frecuencia_recoleccion
    ("frecuencia_recoleccion", "diario", "Diario"),
    ("frecuencia_recoleccion", "dos_tres_semana", "2 a 3 veces por semana"),
    ("frecuencia_recoleccion", "una_semana", "1 vez por semana"),
    ("frecuencia_recoleccion", "irregular_quema", "Irregular / No ingresa recolector (quema o entierro)"),
    
    # estado_vial
    ("estado_vial", "muy_bueno", "Muy bueno (asfaltado óptimo en eje principal E585)"),
    ("estado_vial", "regular", "Regular (vías internas lastradas con baches)"),
    ("estado_vial", "malo_invierno", "Malo (caminos hacia recintos intransitables en invierno)"),
    
    # conectividad_internet
    ("conectividad_internet", "fibra_optica", "Fibra óptica a domicilio"),
    ("conectividad_internet", "datos_moviles", "Datos móviles (3G / 4G)"),
    ("conectividad_internet", "infocentro", "Punto Wi-Fi comunal / Infocentro"),
    ("conectividad_internet", "sin_acceso", "Sin acceso a conectividad"),
    
    # nivel_resolucion_salud
    ("nivel_resolucion_salud", "adecuado", "Adecuado para atención primaria rural"),
    ("nivel_resolucion_salud", "regular", "Regular: Frecuente falta de medicamentos básicos"),
    ("nivel_resolucion_salud", "derivacion_pinas", "Deficiente: Exige traslado urgente al Hospital de Piñas"),
    
    # actividad_economica
    ("actividad_economica", "cana_panela", "Caña de azúcar y molienda de panela"),
    ("actividad_economica", "maiz_granos", "Producción de maíz duro y granos"),
    ("actividad_economica", "platano_cacao", "Plátano y cacao fino de aroma"),
    ("actividad_economica", "ganaderia_bovina", "Ganadería bovina de carne y leche"),
    ("actividad_economica", "porcicultura_avicultura", "Porcicultura y avicultura"),
    ("actividad_economica", "comercio_gastronomia_la_y", "Comercio gastronómico, paraderos y parrilladas en La Y"),
    ("actividad_economica", "transporte_servicios", "Transporte interprovincial, talleres y servicios"),
    
    # canal_comercializacion
    ("canal_comercializacion", "parador_la_y", "Venta directa en paradores gastronómicos de La Y"),
    ("canal_comercializacion", "mercado_pinas_balsas", "Comercialización en mercados de Piñas, Balsas o Machala"),
    ("canal_comercializacion", "intermediario_finca", "Intermediario local a pie de finca"),
    ("canal_comercializacion", "autoconsumo", "Autoconsumo y venta local menor"),
    
    # prioridad_obra
    ("prioridad_obra", "alcantarillado_ptar", "Alcantarillado sanitario y Planta de Tratamiento de Aguas Residuales (PTAR)"),
    ("prioridad_obra", "agua_potable_optimizacion", "Optimización y ampliación del sistema de agua potable y reservorios"),
    ("prioridad_obra", "asfaltado_caminos_recintos", "Asfaltado y mantenimiento de vías hacia recintos productivos (Platanillos, Panupali)"),
    ("prioridad_obra", "ordenamiento_la_y", "Ordenamiento urbano, drenaje y parador turístico en La Y"),
    ("prioridad_obra", "salud_ampliacion", "Ampliación y equipamiento del Centro de Salud Rural"),
    ("prioridad_obra", "defensas_riberenas", "Obras de protección ribereña y mitigación frente a crecidas del río Saracay")
]

for row in choices_data:
    ws_choices.append(list(row))

# 3. SURVEY SHEET
survey_headers = ["type", "name", "label", "hint", "required", "relevant", "appearance"]
ws_survey.append(survey_headers)

survey_rows = [
    # METADATOS Y CONTROL
    ("begin_group", "grp_control", "1. Metadatos de la Encuesta y Ubicación Territorial", "", "", "", ""),
    ("text", "encuestador", "Nombre del Investigador / Encuestador UNL", "Erick Condoy, David Pérez, Camila Correa, Luis Hurtado", "yes", "", ""),
    ("date", "fecha_encuesta", "Fecha de Levantamiento de Campo", "", "yes", "", ""),
    ("geopoint", "coordenadas_gps", "Georreferenciación GPS del punto de entrevista", "Capture ubicación con precisión de teléfono móvil (<10m)", "no", "", ""),
    ("select_one sitio_sector", "sector_sitio", "Sector, Ciudadela o Recinto de Saracay (Piñas, El Oro)", "", "yes", "", ""),
    ("text", "otro_sitio_nom", "Especifique el nombre del recinto o sector", "", "yes", "${sector_sitio} = 'otro_recinto'", ""),
    ("end_group", "grp_control", "", "", "", "", ""),

    # PERFIL DEL INFORMANTE
    ("begin_group", "grp_informante", "2. Perfil del Informante Clave", "", "", "", ""),
    ("select_one cargo_rol", "cargo_rol", "Rol, función u ocupación en la parroquia Saracay", "", "yes", "", ""),
    ("text", "cargo_otro_nom", "Especifique su ocupación o actividad", "", "yes", "${cargo_rol} = 'otro'", ""),
    ("select_one rango_edad", "edad_informante", "Rango de edad del informante", "", "yes", "", ""),
    ("integer", "tiempo_residencia", "¿Cuántos años vive en la Parroquia Saracay?", "Años de residencia continua", "yes", "", ""),
    ("end_group", "grp_informante", "", "", "", "", ""),

    # BLOQUE I: SERVICIOS BASICOS Y AGUA
    ("begin_group", "grp_servicios", "3. Bloque I: Agua, Saneamiento y Residuos", "", "", "", ""),
    ("select_one calidad_agua", "calidad_agua_percepcion", "¿Cómo califica el servicio y calidad del agua en su sector?", "", "yes", "", ""),
    ("select_one continuidad_agua", "continuidad_agua", "Continuidad del suministro de agua durante el día", "", "yes", "", ""),
    ("select_one saneamiento_vivienda", "sistema_excretas", "Tipo de sistema para eliminación de excretas en la vivienda", "", "yes", "", ""),
    ("select_one frecuencia_recoleccion", "recoleccion_basura", "Frecuencia del servicio de recolección de desechos sólidos (GAD Piñas)", "", "yes", "", ""),
    ("text", "problemas_servicios_obs", "Principales deficiencias del servicio de agua / saneamiento", "Ej. turbiedad en invierno, cortes, falta de PTAR, vertidos al río", "no", "", ""),
    ("end_group", "grp_servicios", "", "", "", "", ""),

    # BLOQUE II: VIALIDAD Y CONECTIVIDAD
    ("begin_group", "grp_vialidad", "4. Bloque II: Vialidad, Conectividad y Drenaje", "", "", "", ""),
    ("select_one estado_vial", "estado_vias_acceso", "Estado de las vías secundarias y caminos vecinales a su recinto o finca", "", "yes", "", ""),
    ("select_one conectividad_internet", "acceso_internet", "Condición de acceso a internet en su hogar", "", "yes", "", ""),
    ("select_one si_no", "senal_celular", "¿Dispone de señal celular estable para llamadas de voz?", "", "yes", "", ""),
    ("select_one si_no", "afectacion_inundaciones", "¿Su vivienda o cultivos sufren anegamientos, crecidas o deslizamientos en temporada lluviosa?", "", "yes", "", ""),
    ("end_group", "grp_vialidad", "", "", "", "", ""),

    # BLOQUE III: SALUD Y EDUCACION
    ("begin_group", "grp_salud_educacion", "5. Bloque III: Salud y Equipamiento Comunitario", "", "", "", ""),
    ("select_one nivel_resolucion_salud", "resolucion_salud", "Capacidad de atención del Centro de Salud de Saracay", "", "yes", "", ""),
    ("integer", "tiempo_traslado_pinas", "Tiempo aproximado de traslado al Hospital de Piñas ante emergencias graves", "Minutos en vehículo", "yes", "", ""),
    ("select_one si_no", "infraestructura_escolar_ok", "¿Considera que las instalaciones educativas de su comunidad se encuentran en buen estado?", "", "yes", "", ""),
    ("end_group", "grp_salud_educacion", "", "", "", "", ""),

    # BLOQUE IV: SISTEMA PRODUCTIVO
    ("begin_group", "grp_economia", "6. Bloque IV: Sistema Productivo y Medios de Vida", "", "", "", ""),
    ("select_multiple actividad_economica", "actividades_productivas", "Principales actividades generadoras de ingresos en su hogar (Selección múltiple)", "", "yes", "", ""),
    ("select_one canal_comercializacion", "canal_venta", "Canal predominante de comercialización de su producción", "", "yes", "", ""),
    ("select_one si_no", "asociacion_activa", "¿Forma parte de alguna asociación comunitaria o gremial activa?", "", "yes", "", ""),
    ("text", "trabas_productivas", "Mayor obstáculo para la producción o comercialización en Saracay", "Ej. costos de insumos, precios bajos, vías deterioradas, falta de apoyo técnico", "no", "", ""),
    ("end_group", "grp_economia", "", "", "", "", ""),

    # BLOQUE V: PRIORIDADES Y CARTERA DE INVERSION
    ("begin_group", "grp_prioridades", "7. Bloque V: Prioridades de Inversión y Proyectos Locales", "", "", "", ""),
    ("select_one prioridad_obra", "prioridad_1_inversion", "PRIORIDAD #1: Proyecto u obra pública más urgente para Saracay", "", "yes", "", ""),
    ("select_one prioridad_obra", "prioridad_2_inversion", "PRIORIDAD #2: Segunda obra prioritaria para su comunidad", "", "yes", "", ""),
    ("text", "proyectos_en_ejecucion_conocidos", "¿Conoce usted qué obras o proyectos públicos están actualmente ejecutándose en la parroquia?", "Describa obras del GAD Parroquial, Alcaldía de Piñas o Prefectura de El Oro", "no", "", ""),
    ("image", "foto_respaldo_1", "Fotografía de respaldo 1 (Infraestructura, sector o reunión)", "Capture evidencia en campo", "no", "", ""),
    ("note", "nota_agradecimiento", "Agradecimiento: Muchas gracias por su colaboración con la Universidad Nacional de Loja y el Proyecto de Diagnóstico Territorial de Saracay (Piñas, El Oro).", "", "", "", ""),
    ("end_group", "grp_prioridades", "", "", "", "", "")
]

for row in survey_rows:
    ws_survey.append(list(row))

# Formatting
for sheet in [ws_survey, ws_choices, ws_settings]:
    for col_idx in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

out_path = r"d:\07_Ciclo\septimo_ciclo\public_investment_projects\docs\vaults\proyecto-saracay\levantamiento_de_informacion\encuesta_diagnostico_saracay_kobotoolbox.xlsx"
wb.save(out_path)
print("XLSForm generado exitosamente para Saracay, Piñas, El Oro en:", out_path)
