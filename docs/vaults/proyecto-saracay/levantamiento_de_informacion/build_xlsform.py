# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

# Setup sheets
ws_survey = wb.active
ws_survey.title = 'survey'
ws_choices = wb.create_sheet('choices')
ws_settings = wb.create_sheet('settings')

# Styles
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# 1. SETTINGS SHEET
settings_headers = ['form_title', 'form_id', 'version', 'default_language']
ws_settings.append(settings_headers)
ws_settings.append([
    'UNL - Diagnóstico Territorial y Línea Base de Inversión Pública (Saracay)',
    'unl_diagnostico_saracay_2026',
    '2026101501',
    'Spanish (es)'
])

# 2. CHOICES SHEET
choices_headers = ['list_name', 'name', 'label']
ws_choices.append(choices_headers)

choices_data = [
    # list_cargo
    ('list_cargo', 'autoridad_local', 'Autoridad Local (Presidente GAD - Sr. Jorge Homero Ochoa Sánchez)'),
    ('list_cargo', 'tecnico_gad', 'Técnico/a de Planificación o Infraestructura GAD'),
    ('list_cargo', 'lider_comunal', 'Líder Comunal / Presidente Comunitario'),
    ('list_cargo', 'dirigente_social', 'Dirigente de Organización Social / Junta de Agua'),
    ('list_cargo', 'otro', 'Otro informante clave'),

    # s_no_act
    ('s_no_act', 'si', 'Sí'),
    ('s_no_act', 'no', 'No'),
    ('s_no_act', 'en_actualizacion', 'En actualización'),

    # s_no
    ('s_no', 'si', 'Sí'),
    ('s_no', 'no', 'No'),

    # list_estado_vial
    ('list_estado_vial', 'bueno', 'Bueno (Asfaltado en óptimas condiciones eje E585)'),
    ('list_estado_vial', 'regular', 'Regular (Afirmada / Lastrada con baches o mantenimiento periódico)'),
    ('list_estado_vial', 'malo_trocha', 'Malo / Trocha de tierra (Crítica en invierno hacia recintos)'),

    # list_dinamica_pob
    ('list_dinamica_pob', 'creciendo', 'Creciendo'),
    ('list_dinamica_pob', 'estable', 'Estable'),
    ('list_dinamica_pob', 'disminuyendo', 'Disminuyendo'),

    # list_potencial_tur
    ('list_potencial_tur', 'alto', 'Alto (Nodo gastronómico La Y y balnearios fluviales de río Saracay)'),
    ('list_potencial_tur', 'medio', 'Medio'),
    ('list_potencial_tur', 'bajo', 'Bajo'),

    # list_sector_inv
    ('list_sector_inv', 'saneamiento', 'Agua potable, saneamiento y PTAR'),
    ('list_sector_inv', 'vialidad', 'Vialidad y conectividad con recintos'),
    ('list_sector_inv', 'fomento_prod', 'Fomento agropecuario y comercial en La Y'),
    ('list_sector_inv', 'salud_edu', 'Salud y educación rural'),
    ('list_sector_inv', 'ambiental', 'Gestión ambiental y protección ribereña de cuencas')
]

for row in choices_data:
    ws_choices.append(list(row))

# 3. SURVEY SHEET
survey_headers = ['type', 'name', 'label', 'hint', 'required']
ws_survey.append(survey_headers)

survey_rows = [
    # Módulo 0
    ('begin_group', 'grp_identificacion', 'Módulo 0: Identificación del Informante y Control', '', ''),
    ('text', 'entrevistador', 'Nombre del entrevistador (Equipo Técnico UNL)', 'E. Condoy, D. Pérez, C. Correa, L. Hurtado', 'yes'),
    ('date', 'fecha_entrevista', 'Fecha de la entrevista', '', 'yes'),
    ('text', 'nombre_entrevistado', 'Nombre completo del entrevistado', '', 'yes'),
    ('select_one list_cargo', 'cargo_entrevistado', 'Cargo o rol principal en el territorio', '', 'yes'),
    ('text', 'contacto_entrevistado', 'Teléfono institucional / Correo electrónico', '', 'no'),
    ('end_group', 'grp_identificacion', '', '', ''),

    # Módulo I
    ('begin_group', 'grp_territorio', 'Módulo I: Caracterización General y Territorio', '', ''),
    ('note', 'nota_saracay', 'Jurisdicción: Parroquia Saracay, Cantón Piñas, Provincia de El Oro (DPA: 071056). Superficie: 110,00 km².', '', ''),
    ('select_one s_no_act', 'pdot_vigente', '¿Cuenta la parroquia con PDOT vigente o en proceso de actualización formal?', '', 'yes'),
    ('select_one list_estado_vial', 'vialidad_principal', 'Estado de la conectividad vial principal (Eje E585 - La Y de Saracay)', '', 'yes'),
    ('select_one list_estado_vial', 'vialidad_secundaria', 'Estado general de las vías secundarias y de acceso a los recintos rurales', '', 'yes'),
    ('end_group', 'grp_territorio', '', '', ''),

    # Módulo II
    ('begin_group', 'grp_sociocultural', 'Módulo II: Componente Sociocultural y Servicios Básicos', '', ''),
    ('select_one list_dinamica_pob', 'tendencia_demografica', 'Percepción de la dinámica poblacional en los últimos 5 años', '', 'yes'),
    ('select_one s_no', 'deficiencia_agua', '¿Existen problemas críticos en la continuidad o calidad del servicio de agua de consumo?', '', 'yes'),
    ('select_one s_no', 'deficiencia_alcantarillado', '¿Existe carencia o déficit en el sistema de alcantarillado sanitario y PTAR?', '', 'yes'),
    ('text', 'problemas_salud_frecuentes', 'Principales patologías o problemas de salud atendidos en el Centro de Salud local', 'Ej. parasitosis, diarreas agudas, infecciones respiratorias', 'no'),
    ('end_group', 'grp_sociocultural', '', '', ''),

    # Módulo III
    ('begin_group', 'grp_productivo', 'Módulo III: Componente Económico-Productivo y Ambiental', '', ''),
    ('text', 'principales_cultivos', 'Principales cultivos y rubros productivos de Saracay', 'Caña, maíz, plátano, cacao, cítricos, ganadería bovina', 'yes'),
    ('select_one list_potencial_tur', 'turismo_potencial', 'Potencial turístico y comercial de la parroquia (La Y de Saracay / Balnearios fluviales)', '', 'yes'),
    ('text', 'fuentes_hidricas_clave', 'Ríos o microcuencas prioritarias para consumo y riego', 'Río Saracay, Río Fátima, Río Naranjos', 'no'),
    ('end_group', 'grp_productivo', '', '', ''),

    # Módulo IV
    ('begin_group', 'grp_inversion', 'Módulo IV: Demanda de Inversión Pública', '', ''),
    ('select_one list_sector_inv', 'prioridad_inversion_1', 'Sector más urgente que requiere inversión pública inmediata en Saracay', '', 'yes'),
    ('text', 'proyecto_propuesto_nombre', 'Nombre o idea de proyecto prioritario identificado', 'Ej. Mejoramiento del saneamiento básico y PTAR en Saracay', 'yes'),
    ('text', 'observaciones_adicionales', 'Observaciones finales o recomendaciones técnicas del informante', '', 'no'),
    ('end_group', 'grp_inversion', '', '', '')
]

for row in survey_rows:
    ws_survey.append(list(row))

# Auto width & styles
for ws in [ws_survey, ws_choices, ws_settings]:
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            cell.border = thin_border
            if cell.row == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.font = Font(name='Calibri', size=10)
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

out_file = r'd:\07_Ciclo\septimo_ciclo\public_investment_projects\docs\vaults\proyecto-saracay\levantamiento_de_informacion\encuesta_diagnostico_saracay_kobo_v2.xlsx'
wb.save(out_file)
print('XLSForm V2 guardado exitosamente en:', out_file)
